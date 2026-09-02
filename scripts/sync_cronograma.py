"""Sincroniza o cronograma de cargas ("Compras BD.xlsx", pasta de rede) com o
Supabase. Rode manualmente pra testar, e agende no Agendador de Tarefas do
Windows pra rodar sozinho todo dia de madrugada.

Config: usa o mesmo .env da raiz do projeto (o mesmo do dev_server.py),
nunca commitado — veja .env.example para o formato completo.

Uso: python sync_cronograma.py
"""
import datetime
import hashlib
import json
import logging
import os
import sys
import urllib.error
import urllib.parse
import urllib.request

try:
    import openpyxl
except ImportError:
    print("Falta instalar a dependência: pip install openpyxl")
    sys.exit(1)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(SCRIPT_DIR)  # raiz do projeto — usa o mesmo .env do dev_server.py
CD_VALIDOS = {"GO", "MA", "PA"}
COLUNAS_OBRIGATORIAS = ["CD", "DATA", "DPTO", "FORNECEDOR", "NF, CARGA OU ROMANEIO"]

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    handlers=[
        logging.FileHandler(os.path.join(BASE_DIR, "sync_cronograma.log"), encoding="utf-8"),
        logging.StreamHandler(),
    ],
)
log = logging.getLogger("sync_cronograma")


def load_env_file(path):
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key, value = key.strip(), value.strip()
            if key and key not in os.environ:
                os.environ[key] = value


def norm_codigo(v):
    """Normaliza um valor de célula (int/float/str) pra comparar como texto,
    igual ao comportamento de chave de objeto do JavaScript (que sempre
    converte pra string)."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def make_row_id(cd, data_iso, fornecedor, carga):
    """Precisa bater exatamente com makeRowIdJS em index.html, senão o
    histórico de status vira órfão a cada sync."""
    chave = f"{cd}|{data_iso}|{fornecedor}|{carga}".lower()
    return hashlib.sha1(chave.encode("utf-8")).hexdigest()[:12]


def build_dept_map(wb):
    if "Dados" not in wb.sheetnames:
        return {}
    rows = list(wb["Dados"].iter_rows(values_only=True))
    if not rows:
        return {}
    header = [("" if h is None else str(h).strip()) for h in rows[0]]
    if "Departamento" not in header or "Nome Departamento" not in header:
        return {}
    i_cod, i_nome = header.index("Departamento"), header.index("Nome Departamento")
    mapa = {}
    for row in rows[1:]:
        cod = row[i_cod] if i_cod < len(row) else None
        nome = row[i_nome] if i_nome < len(row) else None
        cod_norm = norm_codigo(cod)
        if cod_norm is not None and nome is not None and cod_norm not in mapa:
            mapa[cod_norm] = str(nome).strip()
    return mapa


def parse_planilha(caminho, mes_inicio):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    if "Retorno Email" not in wb.sheetnames:
        raise RuntimeError('Aba "Retorno Email" não encontrada na planilha.')
    ws = wb["Retorno Email"]
    dept_map = build_dept_map(wb)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise RuntimeError('Aba "Retorno Email" está vazia.')

    header = [("" if h is None else str(h).strip()) for h in rows[0]]
    for col in COLUNAS_OBRIGATORIAS:
        if col not in header:
            raise RuntimeError(f'Coluna obrigatória não encontrada: "{col}"')
    i_cd, i_data = header.index("CD"), header.index("DATA")
    i_dpto, i_forn = header.index("DPTO"), header.index("FORNECEDOR")
    i_carga = header.index("NF, CARGA OU ROMANEIO")

    cargas, ignoradas_cd = [], {}
    ignoradas_sem_data = ignoradas_fora_periodo = 0

    for row in rows[1:]:
        if row is None:
            continue
        cd_raw = row[i_cd] if i_cd < len(row) else None
        data_raw = row[i_data] if i_data < len(row) else None
        if cd_raw is None and data_raw is None:
            continue
        cd = "" if cd_raw is None else str(cd_raw).strip().upper()
        if cd not in CD_VALIDOS:
            chave = cd or "(vazio)"
            ignoradas_cd[chave] = ignoradas_cd.get(chave, 0) + 1
            continue

        if isinstance(data_raw, datetime.datetime):
            data = data_raw.date()
        elif isinstance(data_raw, datetime.date):
            data = data_raw
        else:
            ignoradas_sem_data += 1
            continue
        if data < mes_inicio:
            ignoradas_fora_periodo += 1
            continue

        dpto_cod_norm = norm_codigo(row[i_dpto] if i_dpto < len(row) else None)
        if dpto_cod_norm is not None:
            dpto_nome = dept_map.get(dpto_cod_norm, f"Depto {dpto_cod_norm}")
        else:
            dpto_nome = "Não informado"

        forn_raw = row[i_forn] if i_forn < len(row) else None
        fornecedor = (str(forn_raw).strip() if forn_raw is not None else "") or "Fornecedor não informado"
        carga_raw = row[i_carga] if i_carga < len(row) else None
        carga_txt = (str(carga_raw).strip() if carga_raw is not None else "") or "-"
        data_iso = data.isoformat()

        cargas.append({
            "id": make_row_id(cd, data_iso, fornecedor, carga_txt),
            "cd": cd,
            "data": data_iso,
            "dpto": dpto_nome,
            "fornecedor": fornecedor,
            "carga": carga_txt,
        })

    log.info(
        "Planilha lida: %d carga(s) válidas, %d ignorada(s) sem data, %d fora do período, CDs ignorados: %s",
        len(cargas), ignoradas_sem_data, ignoradas_fora_periodo, ignoradas_cd,
    )
    return cargas


def _supabase_request(method, url, service_role_key, body=None, extra_headers=None):
    headers = {
        "apikey": service_role_key,
        "Authorization": "Bearer " + service_role_key,
        "Content-Type": "application/json",
    }
    if extra_headers:
        headers.update(extra_headers)
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(url, data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            raw = resp.read()
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"Supabase HTTP {e.code} em {method} {url}: {e.read().decode('utf-8', 'replace')}")
    except urllib.error.URLError as e:
        raise RuntimeError(f"Supabase indisponível: {e}")
    return json.loads(raw.decode("utf-8")) if raw else None


def sync_supabase(cargas, supabase_url, service_role_key):
    """Insere cargas novas e atualiza (PATCH — só os campos enviados, nunca
    o status) as já existentes. Não usa upsert por conflito porque o
    comportamento do PostgREST ao omitir colunas num upsert não é garantido
    pela documentação; PATCH é inequivocamente uma atualização parcial."""
    base = supabase_url.rstrip("/")
    if not cargas:
        log.info("Nenhuma carga dentro do período para sincronizar.")
        return

    ids_planilha = [c["id"] for c in cargas]
    ids_existentes = set()
    # busca em lotes pra não estourar limite de tamanho de URL/query
    for i in range(0, len(ids_planilha), 200):
        lote = ids_planilha[i:i + 200]
        filtro = "id=in.(" + ",".join(urllib.parse.quote(x, safe="") for x in lote) + ")"
        encontrados = _supabase_request("GET", f"{base}/rest/v1/cargas?select=id&{filtro}", service_role_key) or []
        ids_existentes.update(r["id"] for r in encontrados)

    agora = datetime.datetime.utcnow().isoformat() + "Z"
    novas, existentes = [], []
    for c in cargas:
        registro = dict(c, atualizado_em=agora)
        (existentes if c["id"] in ids_existentes else novas).append(registro)

    if novas:
        _supabase_request(
            "POST", f"{base}/rest/v1/cargas", service_role_key,
            body=novas, extra_headers={"Prefer": "return=minimal"},
        )
        log.info("%d carga(s) nova(s) inserida(s).", len(novas))

    campos_atualizaveis = ("cd", "data", "dpto", "fornecedor", "carga", "atualizado_em")
    for registro in existentes:
        corpo = {k: registro[k] for k in campos_atualizaveis}
        filtro = "id=eq." + urllib.parse.quote(registro["id"], safe="")
        _supabase_request(
            "PATCH", f"{base}/rest/v1/cargas?{filtro}", service_role_key,
            body=corpo, extra_headers={"Prefer": "return=minimal"},
        )
    if existentes:
        log.info("%d carga(s) existente(s) atualizada(s) (agenda apenas, status preservado).", len(existentes))


def main():
    load_env_file(os.path.join(BASE_DIR, ".env"))
    supabase_url = os.environ.get("SUPABASE_URL")
    service_role_key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    caminho_planilha = os.environ.get("CAMINHO_PLANILHA")
    if not supabase_url or not service_role_key or not caminho_planilha:
        log.error("Configure SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY e CAMINHO_PLANILHA no .env (veja .env.example).")
        sys.exit(1)

    hoje = datetime.date.today()
    mes_inicio = hoje.replace(day=1)

    try:
        cargas = parse_planilha(caminho_planilha, mes_inicio)
        sync_supabase(cargas, supabase_url, service_role_key)
    except Exception:
        log.exception("Falha ao sincronizar o cronograma.")
        sys.exit(1)


if __name__ == "__main__":
    main()
